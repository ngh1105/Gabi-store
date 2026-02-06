"""
Script để chuyển đổi markdown presentation sang PowerPoint
Author: Gabi-store
"""

import re
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def parse_markdown_slides(md_file):
    """Đọc file markdown và tách thành các slides"""
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Tách nội dung theo separator ---
    slides_raw = content.split('\n---\n')
    
    slides = []
    for slide_content in slides_raw:
        if slide_content.strip():
            slides.append(slide_content.strip())
    
    return slides

def extract_title_and_content(slide_text):
    """Trích xuất title và nội dung từ text của slide"""
    lines = slide_text.split('\n')
    
    # Tìm title (dòng đầu tiên bắt đầu bằng #)
    title = ""
    content_lines = []
    title_found = False
    
    for line in lines:
        # Title có thể là # hoặc ##
        if line.startswith('# ') and not title_found:
            title = line.replace('# ', '').strip()
            title_found = True
        elif line.startswith('## ') and not title_found:
            title = line.replace('## ', '').strip()
            title_found = True
        else:
            content_lines.append(line)
    
    content = '\n'.join(content_lines).strip()
    return title, content

def extract_image_path(text):
    """Trích xuất đường dẫn hình ảnh từ markdown"""
    # Pattern: ![alt](path)
    pattern = r'!\[.*?\]\((.*?)\)'
    matches = re.findall(pattern, text)
    return matches

def create_presentation(slides_data, output_file):
    """Tạo PowerPoint presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    for slide_idx, slide_text in enumerate(slides_data):
        print(f"Processing slide {slide_idx + 1}/{len(slides_data)}...")
        
        title, content = extract_title_and_content(slide_text)
        
        # Kiểm tra xem có hình ảnh không
        image_paths = extract_image_path(slide_text)
        
        # Chọn layout
        if image_paths:
            # Slide với hình ảnh - dùng blank layout
            slide_layout = prs.slide_layouts[6]  # Blank layout
        else:
            # Slide text only - dùng title and content
            slide_layout = prs.slide_layouts[1]  # Title and Content
        
        slide = prs.slides.add_slide(slide_layout)
        
        # Thêm title
        if slide_layout == prs.slide_layouts[6]:  # Blank layout
            # Tự tạo title textbox
            title_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.3), Inches(9), Inches(0.8)
            )
            title_frame = title_box.text_frame
            title_frame.text = title
            title_para = title_frame.paragraphs[0]
            title_para.font.size = Pt(32)
            title_para.font.bold = True
            title_para.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue
            title_para.alignment = PP_ALIGN.CENTER
        else:
            slide.shapes.title.text = title
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
            slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
        
        # Xử lý nội dung
        if image_paths:
            # Thêm hình ảnh
            img_top = Inches(1.5)
            for img_path in image_paths:
                if os.path.exists(img_path):
                    try:
                        # Center image
                        slide.shapes.add_picture(
                            img_path, 
                            Inches(1.5), img_top,
                            width=Inches(7)
                        )
                        img_top += Inches(4)  # Offset cho ảnh tiếp theo
                    except Exception as e:
                        print(f"Error adding image {img_path}: {e}")
            
            # Thêm text content bên dưới ảnh (nếu có)
            content_no_images = re.sub(r'!\[.*?\]\(.*?\)', '', content).strip()
            if content_no_images:
                content_box = slide.shapes.add_textbox(
                    Inches(0.5), Inches(5.5), Inches(9), Inches(1.5)
                )
                text_frame = content_box.text_frame
                text_frame.word_wrap = True
                
                # Parse content
                add_formatted_text(text_frame, content_no_images)
        
        else:
            # Slide chỉ có text
            if len(slide.shapes) > 1:
                text_frame = slide.shapes.placeholders[1].text_frame
                text_frame.word_wrap = True
                
                # Parse và thêm formatted text
                add_formatted_text(text_frame, content)
    
    # Lưu file
    prs.save(output_file)
    print(f"\n✅ PowerPoint saved: {output_file}")

def add_formatted_text(text_frame, content):
    """Thêm text với formatting vào text frame"""
    text_frame.clear()
    
    lines = content.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        p = text_frame.add_paragraph()
        
        # Parse markdown formatting
        # Bold: **text**
        # Bullet: - hoặc *
        # Code: `code`
        # Headers: ###
        
        # Check if bullet point
        if line.startswith('- ') or line.startswith('* '):
            line = line[2:]
            p.level = 0
        elif line.startswith('  - ') or line.startswith('  * '):
            line = line[4:]
            p.level = 1
        
        # Check if header
        if line.startswith('### '):
            line = line.replace('### ', '')
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 102, 204)
        elif line.startswith('## '):
            line = line.replace('## ', '')
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 51, 102)
        else:
            p.font.size = Pt(14)
        
        # Remove emojis and special markdown
        line = re.sub(r':[a-z_]+:', '', line)  # Remove emoji codes
        
        # Clean up inline code and bold
        line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)  # Bold
        line = re.sub(r'`(.*?)`', r'\1', line)  # Code
        
        p.text = line

def main():
    # Paths
    script_dir = r"C:\Users\Admin\.gemini\antigravity\brain\8f559d93-0ca4-4178-915c-347e12f2a949"
    
    # Input markdown file
    md_file = os.path.join(
        script_dir,
        'gabi-store-presentation.md'
    )
    
    # Output PowerPoint file
    output_file = r"e:\gabi-store\gabi-store-presentation.pptx"
    
    print("🚀 Starting conversion...")
    print(f"📄 Input:  {md_file}")
    print(f"💾 Output: {output_file}")
    print()
    
    # Parse slides
    slides = parse_markdown_slides(md_file)
    print(f"📊 Found {len(slides)} slides\n")
    
    # Create presentation
    create_presentation(slides, output_file)
    
    print("\n✨ Conversion complete!")
    print(f"📂 File saved at: {output_file}")

if __name__ == "__main__":
    main()
